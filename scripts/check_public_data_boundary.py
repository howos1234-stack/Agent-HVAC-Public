"""Reject restricted product evidence from the public repository."""

from __future__ import annotations

import hashlib
import subprocess
from pathlib import Path

from openpyxl import load_workbook

TEMPLATE = Path("docs/product_data/templates/Agent-HVAC_ProductData_0.2.0_template.xlsx")
EXPECTED_TEMPLATE_SHA256 = "6043d7278a39a4ad5e482ddeb6df54f86fa0ac12e48bd3f5d654fa57d7aeb2cd"
EXPECTED_SHEETS = (
    "products",
    "source_metadata",
    "data_sources",
    "derived_source_conditions",
    "rated_points",
    "rated_point_values",
    "performance_maps",
    "map_axes",
    "performance_points",
    "performance_point_values",
    "map_fixed_conditions",
    "operating_envelopes",
    "envelope_axes",
    "envelope_vertices",
    "envelope_fixed_conditions",
)
EXPECTED_HEADERS = {
    "products": (
        "product_id",
        "component_type",
        "manufacturer",
        "model",
        "supported_refrigerants",
        "status",
        "is_mock",
    ),
    "source_metadata": (
        "product_id",
        "document_ref",
        "original_manufacturer",
        "original_model",
        "retrieved_at",
    ),
    "data_sources": (
        "product_id",
        "source_id",
        "origin",
        "document_ref",
        "table_or_figure_ref",
        "original_manufacturer",
        "original_model",
        "method_id",
        "method_version",
        "equation_ref",
        "input_value_refs",
        "assumptions",
    ),
    "derived_source_conditions": (
        "derived_source_id",
        "value_id",
        "canonical_name",
        "original_name",
        "original_value",
        "original_unit",
        "provided_value_si",
        "provided_canonical_unit",
        "tolerance_kind",
        "tolerance_value",
        "tolerance_unit",
        "tolerance_reference",
        "tolerance_reference_value",
        "tolerance_reference_unit",
        "tolerance_source_id",
        "tolerance_absence_reason",
        "accuracy_validation",
        "source_id",
    ),
    "rated_points": (
        "product_id",
        "rated_point_id",
        "refrigerant_label",
        "topology",
        "mode",
        "source_id",
        "automatic_selection",
        "restriction_reasons",
    ),
    "rated_point_values": (
        "rated_point_id",
        "value_role",
        "value_id",
        "canonical_name",
        "original_name",
        "original_value",
        "original_unit",
        "provided_value_si",
        "provided_canonical_unit",
        "tolerance_kind",
        "tolerance_value",
        "tolerance_unit",
        "tolerance_reference",
        "tolerance_reference_value",
        "tolerance_reference_unit",
        "tolerance_source_id",
        "tolerance_absence_reason",
        "accuracy_validation",
        "source_id",
    ),
    "performance_maps": (
        "product_id",
        "map_id",
        "map_kind",
        "refrigerant_label",
        "topology",
        "mode",
        "interpolation",
        "extrapolation",
        "source_id",
        "automatic_selection",
        "restriction_reasons",
    ),
    "map_axes": (
        "map_id",
        "axis_id",
        "order",
        "canonical_name",
        "native_name",
        "role",
        "physical_kind",
        "original_unit",
        "canonical_unit",
        "source_id",
    ),
    "performance_points": ("map_id", "point_id", "source_id"),
    "performance_point_values": (
        "point_id",
        "value_role",
        "axis_id",
        "value_id",
        "canonical_name",
        "original_name",
        "original_value",
        "original_unit",
        "provided_value_si",
        "provided_canonical_unit",
        "tolerance_kind",
        "tolerance_value",
        "tolerance_unit",
        "tolerance_reference",
        "tolerance_reference_value",
        "tolerance_reference_unit",
        "tolerance_source_id",
        "tolerance_absence_reason",
        "accuracy_validation",
        "source_id",
    ),
    "map_fixed_conditions": (
        "map_id",
        "value_id",
        "canonical_name",
        "original_name",
        "original_value",
        "original_unit",
        "provided_value_si",
        "provided_canonical_unit",
        "tolerance_kind",
        "tolerance_value",
        "tolerance_unit",
        "tolerance_reference",
        "tolerance_reference_value",
        "tolerance_reference_unit",
        "tolerance_source_id",
        "tolerance_absence_reason",
        "accuracy_validation",
        "source_id",
    ),
    "operating_envelopes": (
        "product_id",
        "envelope_id",
        "refrigerant_label",
        "boundary_inclusive",
        "source_id",
        "automatic_selection",
        "restriction_reasons",
    ),
    "envelope_axes": (
        "envelope_id",
        "axis_id",
        "order",
        "canonical_name",
        "native_name",
        "role",
        "physical_kind",
        "original_unit",
        "canonical_unit",
        "source_id",
    ),
    "envelope_vertices": (
        "envelope_id",
        "vertex_order",
        "axis_id",
        "original_value",
        "original_unit",
        "provided_value_si",
        "provided_canonical_unit",
        "source_id",
    ),
    "envelope_fixed_conditions": (
        "envelope_id",
        "value_id",
        "canonical_name",
        "original_name",
        "original_value",
        "original_unit",
        "provided_value_si",
        "provided_canonical_unit",
        "tolerance_kind",
        "tolerance_value",
        "tolerance_unit",
        "tolerance_reference",
        "tolerance_reference_value",
        "tolerance_reference_unit",
        "tolerance_source_id",
        "tolerance_absence_reason",
        "accuracy_validation",
        "source_id",
    ),
}
PUBLIC_BINARY_ALLOWLIST = {
    TEMPLATE,
    Path("docs/validation/ws_e/workbench_canvas_closed_loop.png"),
    Path("docs/validation/ws_e/workbench_converged_states.png"),
    Path("docs/validation/ws_e/workbench_unsupported_component.png"),
}
RESTRICTED_SUFFIXES = {
    ".gif",
    ".jpeg",
    ".jpg",
    ".pdf",
    ".png",
    ".webp",
    ".xls",
    ".xlsb",
    ".xlsm",
    ".xlsx",
    ".zip",
}


def repository_files(root: Path) -> list[Path]:
    result = subprocess.run(
        ["git", "-C", str(root), "ls-files", "-z"],
        check=True,
        capture_output=True,
    )
    return [root / Path(item) for item in result.stdout.decode().split("\0") if item]


def check_restricted_artifacts(root: Path, tracked_files: list[Path]) -> list[str]:
    return [
        f"Restricted public artifact: {path.relative_to(root).as_posix()}"
        for path in tracked_files
        if path.suffix.lower() in RESTRICTED_SUFFIXES
        and path.relative_to(root) not in PUBLIC_BINARY_ALLOWLIST
    ]


def check_component_data(root: Path, tracked_files: list[Path]) -> list[str]:
    component_root = root / "data/components"
    if not component_root.exists():
        return []
    return [
        f"Public component data must remain empty: {path.relative_to(root).as_posix()}"
        for path in tracked_files
        if path.is_relative_to(component_root) and path.name != ".gitkeep"
    ]


def check_blank_template(root: Path) -> list[str]:
    path = root / TEMPLATE
    if not path.is_file():
        return [f"Missing public blank template: {TEMPLATE.as_posix()}"]

    errors: list[str] = []
    actual_sha256 = hashlib.sha256(path.read_bytes()).hexdigest()
    if actual_sha256 != EXPECTED_TEMPLATE_SHA256:
        errors.append(
            "Public template SHA-256 differs: "
            f"expected {EXPECTED_TEMPLATE_SHA256}, got {actual_sha256}"
        )

    workbook = load_workbook(path, read_only=False, data_only=False)
    try:
        if tuple(workbook.sheetnames) != EXPECTED_SHEETS:
            errors.append(
                "Public template sheet contract differs: "
                f"expected {list(EXPECTED_SHEETS)!r}, got {workbook.sheetnames!r}"
            )
        for sheet in workbook.worksheets:
            expected_header = EXPECTED_HEADERS.get(sheet.title)
            actual_header = tuple(
                cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))
            )
            if expected_header is not None and actual_header != expected_header:
                errors.append(
                    f"Public template header differs: {sheet.title}!1; "
                    f"expected {expected_header!r}, got {actual_header!r}"
                )
            for row_number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
                if any(value not in (None, "") for value in row):
                    errors.append(f"Public template contains data: {sheet.title}!{row_number}")
        return errors
    finally:
        workbook.close()


def check(root: Path, tracked_files: list[Path] | None = None) -> list[str]:
    files = repository_files(root) if tracked_files is None else tracked_files
    return [
        *check_restricted_artifacts(root, files),
        *check_component_data(root, files),
        *check_blank_template(root),
    ]


def main() -> int:
    root = Path(__file__).resolve().parents[1]
    errors = check(root)
    if errors:
        print("\n".join(errors))
        return 1
    print(
        "Public data boundary verified: one blank 15-sheet template, "
        "no restricted artifacts or component data"
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
