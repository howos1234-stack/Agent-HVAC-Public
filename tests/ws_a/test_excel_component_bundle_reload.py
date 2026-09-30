"""Exercise multi-workbook isolation and reload recovery with synthetic data."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import load_workbook

from agent_hvac.components.heat_exchanger_rated_point import (
    HeatExchangerRatedPointError,
    HeatExchangerRatedPointSelection,
    select_heat_exchanger_rated_point,
)
from agent_hvac.components.product_adapter import adapt_product_performance_map
from agent_hvac.database.index import ExcelComponentRepository
from agent_hvac.schemas.components import (
    ComponentType,
    OperatingMode,
    ProductRecord,
    ProductTopology,
)
from agent_hvac.utils.units import Quantity
from tests.ws_a.test_excel_component_bundle_integration import _write_bundle


def _product_ids(repository: ExcelComponentRepository) -> set[str]:
    return {record.product_id for record in repository.records}


def _replace_product_id(path: Path, replacement: str) -> None:
    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        headers = tuple(cell.value for cell in sheet[1])
        if "product_id" not in headers:
            continue
        column = headers.index("product_id") + 1
        for row in range(2, sheet.max_row + 1):
            if sheet.cell(row=row, column=column).value:
                sheet.cell(row=row, column=column).value = replacement
    workbook.save(path)
    workbook.close()


def _gas_cooler_selection(product: ProductRecord) -> HeatExchangerRatedPointSelection:
    return HeatExchangerRatedPointSelection(
        product=product,
        rated_point_id="gas-cooler-rated-1",
        refrigerant="R744",
        topology=ProductTopology.R744_TRANSCRITICAL,
        operating_mode=OperatingMode.COOLING,
        conditions={
            "inlet_pressure": Quantity(value=90.0, unit="bar"),
            "inlet_temperature": Quantity(value=375.0, unit="K"),
            "mass_flow": Quantity(value=360.0, unit="kg/h"),
        },
    )


def test_reload_removes_failed_product_and_restores_it_after_workbook_repair(
    tmp_path: Path,
) -> None:
    _write_bundle(tmp_path)
    compressor = tmp_path / "synthetic-compressor.xlsx"
    original = compressor.read_bytes()
    repository = ExcelComponentRepository(tmp_path)

    initial = repository.reload()
    unchanged = repository.reload()
    assert initial.database_version == unchanged.database_version
    assert initial.loaded_products == 3

    compressor.write_bytes(b"not an xlsx zip archive")
    failed = repository.reload()
    assert failed.database_version != initial.database_version
    assert failed.rejected_files == 1
    assert _product_ids(repository) == {
        "valve-synthetic-compressor-1",
        "gas-cooler-synthetic-compressor-1",
    }
    assert all(
        source.source_id.startswith(("valve-", "gas-cooler-"))
        for product in repository.records
        for source in product.data_sources
    )
    valve = next(
        product
        for product in repository.records
        if product.component_type == ComponentType.EXPANSION_VALVE
    )
    valve_map = adapt_product_performance_map(valve, "valve-cmp-map-1")
    assert valve_map.source_ids == ("valve-src-cmp-1",)

    compressor.write_bytes(original)
    recovered = repository.reload()
    assert recovered.database_version == initial.database_version
    assert recovered.errors == ()
    assert recovered.loaded_products == 3
    assert "compressor-synthetic-compressor-1" in _product_ids(repository)


def test_duplicate_product_id_rejects_both_owners_and_recovers_without_stale_records(
    tmp_path: Path,
) -> None:
    _write_bundle(tmp_path)
    valve = tmp_path / "synthetic-valve.xlsx"
    original = valve.read_bytes()
    repository = ExcelComponentRepository(tmp_path)
    baseline = repository.reload()

    _replace_product_id(valve, "compressor-synthetic-compressor-1")
    duplicate = repository.reload()
    assert duplicate.database_version != baseline.database_version
    assert duplicate.rejected_files == 2
    assert duplicate.loaded_products == 1
    assert _product_ids(repository) == {"gas-cooler-synthetic-compressor-1"}
    assert any("duplicate product_id" in error for error in duplicate.errors)

    valve.write_bytes(original)
    recovered = repository.reload()
    assert recovered.database_version == baseline.database_version
    assert recovered.errors == ()
    assert recovered.loaded_products == 3


def test_bad_cross_workbook_source_reference_quarantines_only_its_parent_and_recovers(
    tmp_path: Path,
) -> None:
    _write_bundle(tmp_path)
    gas_cooler = tmp_path / "synthetic-gas-cooler.xlsx"
    original = gas_cooler.read_bytes()
    repository = ExcelComponentRepository(tmp_path)
    baseline = repository.reload()

    workbook = load_workbook(gas_cooler)
    values = workbook["rated_point_values"]
    headers = tuple(cell.value for cell in values[1])
    source_column = headers.index("source_id") + 1
    values.cell(row=2, column=source_column).value = "compressor-src-cmp-1"
    workbook.save(gas_cooler)
    workbook.close()

    invalid = repository.reload()
    products = {record.component_type: record for record in repository.records}
    assert invalid.database_version != baseline.database_version
    assert invalid.rejected_files == 0
    assert invalid.loaded_products == 3
    assert len(products[ComponentType.COMPRESSOR].performance_maps) == 1
    assert len(products[ComponentType.EXPANSION_VALVE].performance_maps) == 1
    assert products[ComponentType.GAS_COOLER].rated_points == ()
    assert any("source_id" in error for error in invalid.errors)
    with pytest.raises(HeatExchangerRatedPointError, match="does not contain rated point"):
        select_heat_exchanger_rated_point(_gas_cooler_selection(products[ComponentType.GAS_COOLER]))

    gas_cooler.write_bytes(original)
    recovered = repository.reload()
    assert recovered.database_version == baseline.database_version
    assert recovered.errors == ()
    restored = {record.component_type: record for record in repository.records}
    assert len(restored[ComponentType.GAS_COOLER].rated_points) == 1
    rated = select_heat_exchanger_rated_point(
        _gas_cooler_selection(restored[ComponentType.GAS_COOLER])
    )
    assert rated.rated_point_source_id == "gas-cooler-src-cmp-1"


def test_reload_add_modify_delete_does_not_retain_removed_product_or_source(
    tmp_path: Path,
) -> None:
    _write_bundle(tmp_path)
    valve = tmp_path / "synthetic-valve.xlsx"
    original = valve.read_bytes()
    repository = ExcelComponentRepository(tmp_path)
    baseline = repository.reload()

    valve.unlink()
    deleted = repository.reload()
    assert deleted.database_version != baseline.database_version
    assert deleted.loaded_products == 2
    assert ComponentType.EXPANSION_VALVE not in {
        record.component_type for record in repository.records
    }
    assert all(
        not source.source_id.startswith("valve-")
        for product in repository.records
        for source in product.data_sources
    )

    valve.write_bytes(original)
    restored = repository.reload()
    assert restored.database_version == baseline.database_version
    assert restored.loaded_products == 3

    workbook = load_workbook(valve)
    products = workbook["products"]
    headers = tuple(cell.value for cell in products[1])
    products.cell(row=2, column=headers.index("model") + 1).value = "Synthetic Valve Reloaded"
    workbook.save(valve)
    workbook.close()
    modified = repository.reload()
    assert modified.database_version != restored.database_version
    valve_record = next(
        record
        for record in repository.records
        if record.component_type == ComponentType.EXPANSION_VALVE
    )
    assert valve_record.model == "Synthetic Valve Reloaded"
