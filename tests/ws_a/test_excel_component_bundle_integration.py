"""Load a synthetic compressor, valve and HX bundle through the public contract."""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from agent_hvac.components.heat_exchanger_rated_point import (
    HeatExchangerRatedPointSelection,
    select_heat_exchanger_rated_point,
)
from agent_hvac.components.product_adapter import adapt_product_performance_map
from agent_hvac.database.index import ExcelComponentRepository
from agent_hvac.schemas.components import (
    ComponentType,
    OperatingMode,
    ProductRecord,
    ProductTopology,
)
from agent_hvac.utils.units import Quantity
from tests.ws_a.test_excel_hx_integration import (
    _convert_detailed_workbook_to_gas_cooler,
)
from tests.ws_b.test_excel_repository import (
    convert_detailed_workbook_to_valve,
    write_detailed_workbook,
)


def _namespace_identifiers(path: Path, namespace: str) -> None:
    workbook = load_workbook(path)
    for sheet in workbook.worksheets:
        headers = tuple(cell.value for cell in sheet[1])
        id_columns = {
            index + 1
            for index, header in enumerate(headers)
            if isinstance(header, str) and header.endswith("_id")
        }
        for row in sheet.iter_rows(min_row=2):
            for column in id_columns:
                cell = row[column - 1]
                if isinstance(cell.value, str) and cell.value:
                    cell.value = f"{namespace}-{cell.value}"
    workbook.save(path)
    workbook.close()


def _write_bundle(root: Path) -> None:
    compressor = root / "synthetic-compressor.xlsx"
    valve = root / "synthetic-valve.xlsx"
    gas_cooler = root / "synthetic-gas-cooler.xlsx"

    write_detailed_workbook(compressor)
    _namespace_identifiers(compressor, "compressor")

    write_detailed_workbook(valve)
    convert_detailed_workbook_to_valve(valve)
    _namespace_identifiers(valve, "valve")

    write_detailed_workbook(gas_cooler)
    _convert_detailed_workbook_to_gas_cooler(gas_cooler)
    _namespace_identifiers(gas_cooler, "gas-cooler")


def _record_ids(product: ProductRecord) -> dict[str, set[str]]:
    value_ids: set[str] = set()
    point_ids: set[str] = set()
    for rated_point in product.rated_points:
        value_ids.update(
            value.value_id for value in (*rated_point.conditions, *rated_point.outputs)
        )
    for performance_map in product.performance_maps:
        value_ids.update(value.value_id for value in performance_map.fixed_conditions)
        for point in performance_map.points:
            point_ids.add(point.point_id)
            value_ids.update(value.value_id for value in (*point.coordinates, *point.outputs))
    return {
        "product": {product.product_id},
        "map": {performance_map.map_id for performance_map in product.performance_maps},
        "rated_point": {rated.rated_point_id for rated in product.rated_points},
        "point": point_ids,
        "value": value_ids,
    }


def _assert_product_record_ids_are_disjoint(products: list[ProductRecord]) -> None:
    seen = {kind: set() for kind in ("product", "map", "rated_point", "point", "value")}
    for product in products:
        for kind, identifiers in _record_ids(product).items():
            assert seen[kind].isdisjoint(identifiers), (kind, identifiers & seen[kind])
            seen[kind].update(identifiers)


def test_synthetic_component_bundle_reaches_ws_a_adapters_without_id_collisions(
    tmp_path: Path,
) -> None:
    _write_bundle(tmp_path)
    repository = ExcelComponentRepository(tmp_path)

    summary = repository.reload()

    assert summary.errors == ()
    assert summary.loaded_products == 3
    products = {product.component_type: product for product in repository.records}
    assert set(products) == {
        ComponentType.COMPRESSOR,
        ComponentType.EXPANSION_VALVE,
        ComponentType.GAS_COOLER,
    }
    assert all(product.is_mock for product in products.values())
    assert {product.product_id for product in products.values()} == {
        "compressor-synthetic-compressor-1",
        "valve-synthetic-compressor-1",
        "gas-cooler-synthetic-compressor-1",
    }
    _assert_product_record_ids_are_disjoint(list(products.values()))
    assert (
        len({source.source_id for product in products.values() for source in product.data_sources})
        == 3
    )

    compressor_map = adapt_product_performance_map(
        products[ComponentType.COMPRESSOR],
        "compressor-cmp-map-1",
    )
    valve_map = adapt_product_performance_map(
        products[ComponentType.EXPANSION_VALVE],
        "valve-cmp-map-1",
    )
    assert compressor_map.source_ids == ("compressor-src-cmp-1",)
    assert valve_map.source_ids == ("valve-src-cmp-1",)
    assert all(
        value.original_unit and value.source_id.startswith("compressor-")
        for point in products[ComponentType.COMPRESSOR].performance_maps[0].points
        for value in (*point.coordinates, *point.outputs)
    )
    assert all(
        value.original_unit and value.source_id.startswith("valve-")
        for point in products[ComponentType.EXPANSION_VALVE].performance_maps[0].points
        for value in (*point.coordinates, *point.outputs)
    )

    gas_cooler = products[ComponentType.GAS_COOLER]
    rated = select_heat_exchanger_rated_point(
        HeatExchangerRatedPointSelection(
            product=gas_cooler,
            rated_point_id="gas-cooler-rated-1",
            refrigerant="R744",
            topology=ProductTopology.R744_TRANSCRITICAL,
            operating_mode=OperatingMode.COOLING,
            conditions={
                "inlet_pressure": Quantity(value=90.0, unit="bar"),
                "inlet_temperature": Quantity(value=375.0, unit="K"),
                "mass_flow": Quantity(value=360.0, unit="kg/h"),
            },
        )
    )
    assert rated.rated_point_source_id == "gas-cooler-src-cmp-1"
    assert rated.is_mock is True
    assert all(
        value.original_unit and value.source_id.startswith("gas-cooler-")
        for value in (*rated.rated_conditions, *rated.rated_outputs)
    )
