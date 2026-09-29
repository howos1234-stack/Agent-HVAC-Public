"""End-to-end checks from Excel schema 0.2.0 to the P05 HX calculation."""

from __future__ import annotations

from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from agent_hvac.components.heat_exchanger_1d import (
    HeatExchanger1DInput,
    SpecificHeatCapacity,
    evaluate_heat_exchanger_1d,
)
from agent_hvac.components.heat_exchanger_rated_point import (
    HeatExchangerBoundaryInput,
    HeatExchangerRatedPointError,
    HeatExchangerRatedPointSelection,
    evaluate_heat_exchanger_rated_point,
    select_heat_exchanger_rated_point,
)
from agent_hvac.database.index import ExcelComponentRepository
from agent_hvac.properties.coolprop_backend import CoolPropBackend
from agent_hvac.schemas.components import (
    ComponentType,
    OperatingMode,
    ProductRecord,
    ProductTopology,
)
from agent_hvac.utils.units import MassFlow, Pressure, Quantity, Temperature
from tests.ws_b.test_excel_repository import value_cells, write_detailed_workbook

_DETAIL_SHEETS = (
    "performance_maps",
    "map_axes",
    "performance_points",
    "performance_point_values",
    "map_fixed_conditions",
    "operating_envelopes",
    "envelope_axes",
    "envelope_vertices",
    "envelope_fixed_conditions",
)


def _column(sheet: Any, name: str) -> int:
    headers = list(next(sheet.iter_rows(values_only=True)))
    return headers.index(name) + 1


def _convert_detailed_workbook_to_gas_cooler(path: Path, *, ua_unit: str = "kW/K") -> None:
    workbook = load_workbook(path)
    products = workbook["products"]
    products.cell(row=2, column=_column(products, "component_type")).value = "gas_cooler"

    for sheet_name in _DETAIL_SHEETS:
        sheet = workbook[sheet_name]
        if sheet.max_row > 1:
            sheet.delete_rows(2, sheet.max_row - 1)

    values = workbook["rated_point_values"]
    values.delete_rows(2, values.max_row - 1)
    for value_id, role, name, value, unit in (
        ("hx-inlet-pressure", "CONDITION", "inlet_pressure", 90.0, "bar"),
        ("hx-inlet-temperature", "CONDITION", "inlet_temperature", 375.0, "K"),
        ("hx-mass-flow", "CONDITION", "mass_flow", 360.0, "kg/h"),
        ("hx-ua", "OUTPUT", "ua", 0.2, ua_unit),
        (
            "hx-pressure-drop",
            "OUTPUT",
            "refrigerant_pressure_drop",
            0.2,
            "bar",
        ),
    ):
        values.append(["rated-1", role, *value_cells(value_id, name, value, unit)])

    workbook.save(path)
    workbook.close()


def _conditions() -> dict[str, Quantity]:
    return {
        "inlet_pressure": Quantity(value=90.0, unit="bar"),
        "inlet_temperature": Quantity(value=375.0, unit="K"),
        "mass_flow": Quantity(value=360.0, unit="kg/h"),
    }


def _selection(product: ProductRecord) -> HeatExchangerRatedPointSelection:
    return HeatExchangerRatedPointSelection(
        product=product,
        rated_point_id="rated-1",
        refrigerant="R744",
        topology=ProductTopology.R744_TRANSCRITICAL,
        operating_mode=OperatingMode.COOLING,
        conditions=_conditions(),
    )


def _boundaries() -> HeatExchangerBoundaryInput:
    backend = CoolPropBackend()
    return HeatExchangerBoundaryInput(
        refrigerant_inlet_state=backend.state_pt(
            "R744",
            Pressure(value=9_000_000.0, unit="Pa"),
            Temperature(value=375.0, unit="K"),
        ),
        refrigerant_mass_flow=MassFlow(value=0.1, unit="kg/s"),
        secondary_inlet_temperature=Temperature(value=290.0, unit="K"),
        secondary_mass_flow=MassFlow(value=1.0, unit="kg/s"),
        secondary_specific_heat_capacity=SpecificHeatCapacity(
            value=1_000.0,
            unit="J/(kg*K)",
        ),
        cell_count=80,
    )


def test_excel_loader_output_drives_hx_rated_point_and_matches_direct_p05(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "synthetic-gas-cooler.xlsx"
    write_detailed_workbook(workbook_path)
    _convert_detailed_workbook_to_gas_cooler(workbook_path)
    repository = ExcelComponentRepository(tmp_path)

    summary = repository.reload()

    assert summary.errors == ()
    assert summary.loaded_products == 1
    product = ProductRecord.model_validate(repository.records[0].model_dump(mode="json"))
    assert product.component_type == ComponentType.GAS_COOLER
    assert product.performance_maps == ()
    assert product.operating_envelopes == ()

    rated = select_heat_exchanger_rated_point(_selection(product))
    boundaries = _boundaries()
    backend = CoolPropBackend()
    connected = evaluate_heat_exchanger_rated_point(backend, rated, boundaries)
    direct = evaluate_heat_exchanger_1d(
        backend,
        HeatExchanger1DInput(
            mode=rated.heat_exchanger_mode,
            refrigerant_inlet_state=boundaries.refrigerant_inlet_state,
            refrigerant_mass_flow=boundaries.refrigerant_mass_flow,
            secondary_inlet_temperature=boundaries.secondary_inlet_temperature,
            secondary_mass_flow=boundaries.secondary_mass_flow,
            secondary_specific_heat_capacity=boundaries.secondary_specific_heat_capacity,
            total_thermal_conductance=rated.total_thermal_conductance,
            refrigerant_pressure_drop=rated.refrigerant_pressure_drop,
            cell_count=boundaries.cell_count,
        ),
    )

    assert connected.calculation == direct
    assert connected.model_path == "RATED_POINT_CONSTANT_UA"
    assert rated.product_id == "synthetic-compressor-1"
    assert rated.rated_point_id == "rated-1"
    assert rated.total_thermal_conductance.value == pytest.approx(200.0, abs=0.0)
    assert rated.refrigerant_pressure_drop.value == pytest.approx(20_000.0, abs=0.0)
    assert rated.rated_point_source_id == "src-cmp-1"
    assert rated.is_mock is True

    conditions = {value.canonical_name: value for value in rated.rated_conditions}
    outputs = {value.canonical_name: value for value in rated.rated_outputs}
    assert conditions["inlet_pressure"].original_unit == "bar"
    assert conditions["inlet_pressure"].value == pytest.approx(9_000_000.0, abs=0.0)
    assert conditions["mass_flow"].original_unit == "kg/h"
    assert conditions["mass_flow"].value == pytest.approx(0.1, rel=0.0, abs=1e-15)
    assert outputs["ua"].original_unit == "kW/K"
    assert outputs["ua"].value == pytest.approx(200.0, abs=0.0)
    assert outputs["refrigerant_pressure_drop"].original_unit == "bar"
    assert outputs["refrigerant_pressure_drop"].value == pytest.approx(20_000.0, abs=0.0)
    assert {value.source_id for value in (*rated.rated_conditions, *rated.rated_outputs)} == {
        "src-cmp-1"
    }


def test_quarantined_excel_hx_rated_point_cannot_fall_back_to_default_ua(
    tmp_path: Path,
) -> None:
    workbook_path = tmp_path / "invalid-gas-cooler.xlsx"
    write_detailed_workbook(workbook_path)
    _convert_detailed_workbook_to_gas_cooler(workbook_path, ua_unit="m")
    repository = ExcelComponentRepository(tmp_path)

    summary = repository.reload()

    assert summary.loaded_products == 1
    product = repository.records[0]
    assert product.rated_points == ()
    assert any("incompatible" in error and "rated-1" in error for error in summary.errors)
    with pytest.raises(HeatExchangerRatedPointError, match="does not contain rated point"):
        select_heat_exchanger_rated_point(_selection(product))
