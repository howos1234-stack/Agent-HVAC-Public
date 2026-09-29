"""Synthetic workbook tests for the WS-B P07 minimum loader."""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from agent_hvac.database.detailed_loader import (
    DERIVED_SOURCE_CONDITION_COLUMNS,
    ENVELOPE_AXIS_COLUMNS,
    ENVELOPE_FIXED_COLUMNS,
    ENVELOPE_PARENT_COLUMNS,
    ENVELOPE_VERTEX_COLUMNS,
    MAP_AXIS_COLUMNS,
    MAP_FIXED_COLUMNS,
    MAP_PARENT_COLUMNS,
    POINT_COLUMNS,
    POINT_VALUE_COLUMNS,
    RATED_PARENT_COLUMNS,
    RATED_VALUE_COLUMNS,
)
from agent_hvac.database.detailed_loader import (
    SOURCE_COLUMNS as DETAIL_SOURCE_COLUMNS,
)
from agent_hvac.database.excel_loader import (
    PRODUCT_COLUMNS,
    SOURCE_COLUMNS,
    discover_workbooks,
)
from agent_hvac.database.index import ExcelComponentRepository
from agent_hvac.schemas.components import (
    ComponentQuery,
    ComponentType,
    PhysicsDerivedDataSource,
)

PRODUCT_HEADER = [
    "product_id",
    "component_type",
    "manufacturer",
    "model",
    "supported_refrigerants",
    "status",
    "is_mock",
    "rated_capacity_value",
    "rated_capacity_unit",
]
SOURCE_HEADER = [
    "product_id",
    "document_ref",
    "original_manufacturer",
    "original_model",
    "retrieved_at",
]

TOLERANCE_HEADER = [
    "tolerance_kind",
    "tolerance_value",
    "tolerance_unit",
    "tolerance_reference",
    "tolerance_reference_value",
    "tolerance_reference_unit",
    "tolerance_source_id",
    "tolerance_absence_reason",
    "accuracy_validation",
]
VALUE_HEADER = [
    "value_id",
    "canonical_name",
    "original_name",
    "original_value",
    "original_unit",
    "provided_value_si",
    "provided_canonical_unit",
    *TOLERANCE_HEADER,
    "source_id",
]

TEMPLATE_PATH = (
    Path(__file__).resolve().parents[2]
    / "docs"
    / "product_data"
    / "templates"
    / "Agent-HVAC_ProductData_0.2.0_template.xlsx"
)


def test_excel_020_template_matches_loader_contract() -> None:
    expected = {
        "products": PRODUCT_COLUMNS,
        "source_metadata": SOURCE_COLUMNS,
        "data_sources": DETAIL_SOURCE_COLUMNS,
        "derived_source_conditions": DERIVED_SOURCE_CONDITION_COLUMNS,
        "rated_points": RATED_PARENT_COLUMNS,
        "rated_point_values": RATED_VALUE_COLUMNS,
        "performance_maps": MAP_PARENT_COLUMNS,
        "map_axes": MAP_AXIS_COLUMNS,
        "performance_points": POINT_COLUMNS,
        "performance_point_values": POINT_VALUE_COLUMNS,
        "map_fixed_conditions": MAP_FIXED_COLUMNS,
        "operating_envelopes": ENVELOPE_PARENT_COLUMNS,
        "envelope_axes": ENVELOPE_AXIS_COLUMNS,
        "envelope_vertices": ENVELOPE_VERTEX_COLUMNS,
        "envelope_fixed_conditions": ENVELOPE_FIXED_COLUMNS,
    }

    workbook = load_workbook(TEMPLATE_PATH, read_only=False, data_only=False)
    try:
        assert workbook.sheetnames == list(expected)
        assert all(len(name) <= 31 for name in workbook.sheetnames)
        for sheet_name, columns in expected.items():
            sheet = workbook[sheet_name]
            assert tuple(cell.value for cell in sheet[1]) == columns
            assert sheet.freeze_panes == "A2"
            assert all(cell.data_type != "f" for row in sheet.iter_rows() for cell in row)
    finally:
        workbook.close()


def value_cells(
    value_id: str,
    name: str,
    value: float,
    unit: str,
    *,
    provided: float | None = None,
    provided_unit: str | None = None,
    source_id: str = "src-cmp-1",
) -> list[object]:
    return [
        value_id,
        name,
        name.replace("_", " ").title(),
        value,
        unit,
        provided,
        provided_unit,
        None,
        None,
        None,
        None,
        None,
        None,
        None,
        "NOT_STATED_IN_SOURCE",
        "NOT_VERIFIABLE",
        source_id,
    ]


def add_detailed_compressor_sheets(workbook: Workbook, *, incomplete_map: bool = False) -> None:
    source = workbook.create_sheet("data_sources")
    source.append(
        [
            "product_id",
            "source_id",
            "origin",
            "document_ref",
            "table_or_figure_ref",
            "original_manufacturer",
            "original_model",
            "method_id",
            "method_version",
            "equation_ref",
            "input_value_refs",
            "assumptions",
        ]
    )
    source.append(
        [
            "synthetic-compressor-1",
            "src-cmp-1",
            "MANUFACTURER",
            "synthetic-detail-fixture",
            "Table 1",
            "Synthetic Manufacturer",
            "Mock Model",
            None,
            None,
            None,
            "[]",
            "[]",
        ]
    )

    rated = workbook.create_sheet("rated_points")
    rated.append(
        [
            "product_id",
            "rated_point_id",
            "refrigerant_label",
            "topology",
            "mode",
            "source_id",
            "automatic_selection",
            "restriction_reasons",
        ]
    )
    rated.append(
        [
            "synthetic-compressor-1",
            "rated-1",
            "CO2",
            "R744_TRANSCRITICAL",
            "COOLING",
            "src-cmp-1",
            "ELIGIBLE",
            "[]",
        ]
    )
    rated_values = workbook.create_sheet("rated_point_values")
    rated_values.append(["rated_point_id", "value_role", *VALUE_HEADER])
    for value_id, name, value, unit in (
        ("rated-suction-p", "suction_pressure", 30.0, "bar"),
        ("rated-discharge-p", "discharge_pressure", 90.0, "bar"),
        ("rated-suction-t", "suction_temperature", 6.85, "degC"),
        ("rated-gc-t", "gas_cooler_outlet_temperature", 35.0, "degC"),
        ("rated-frequency", "frequency", 50.0, "Hz"),
    ):
        rated_values.append(["rated-1", "CONDITION", *value_cells(value_id, name, value, unit)])
    rated_values.append(
        ["rated-1", "OUTPUT", *value_cells("rated-power", "input_power", 2.0, "kW")]
    )
    rated_values.append(
        ["rated-1", "OUTPUT", *value_cells("rated-flow", "mass_flow", 180.0, "kg/h")]
    )

    maps = workbook.create_sheet("performance_maps")
    maps.append(
        [
            "product_id",
            "map_id",
            "map_kind",
            "refrigerant_label",
            "topology",
            "mode",
            "interpolation",
            "extrapolation",
            "source_id",
            "automatic_selection",
            "restriction_reasons",
        ]
    )
    maps.append(
        [
            "synthetic-compressor-1",
            "cmp-map-1",
            "COMPRESSOR",
            "CarbonDioxide",
            "R744_TRANSCRITICAL",
            "COOLING",
            "MULTILINEAR",
            "FORBIDDEN",
            "src-cmp-1",
            "ELIGIBLE",
            "[]",
        ]
    )
    axes = workbook.create_sheet("map_axes")
    axes.append(
        [
            "map_id",
            "axis_id",
            "order",
            "canonical_name",
            "native_name",
            "role",
            "physical_kind",
            "original_unit",
            "canonical_unit",
            "source_id",
        ]
    )
    axes.append(
        [
            "cmp-map-1",
            "axis-suction",
            0,
            "suction_pressure",
            "Suction pressure",
            "CANONICAL",
            "PRESSURE",
            "bar",
            "Pa",
            "src-cmp-1",
        ]
    )
    axes.append(
        [
            "cmp-map-1",
            "axis-discharge",
            1,
            "discharge_pressure",
            "Discharge pressure",
            "CANONICAL",
            "PRESSURE",
            "bar",
            "Pa",
            "src-cmp-1",
        ]
    )
    points = workbook.create_sheet("performance_points")
    points.append(["map_id", "point_id", "source_id"])
    point_values = workbook.create_sheet("performance_point_values")
    point_values.append(["point_id", "value_role", "axis_id", *VALUE_HEADER])
    coordinates = [(30.0, 80.0), (30.0, 90.0), (40.0, 80.0), (40.0, 90.0)]
    if incomplete_map:
        coordinates.pop()
    for index, (suction, discharge) in enumerate(coordinates, start=1):
        point_id = f"point-{index}"
        points.append(["cmp-map-1", point_id, "src-cmp-1"])
        point_values.append(
            [
                point_id,
                "COORDINATE",
                "axis-suction",
                *value_cells(f"{point_id}-suction", "suction_pressure", suction, "bar"),
            ]
        )
        point_values.append(
            [
                point_id,
                "COORDINATE",
                "axis-discharge",
                *value_cells(f"{point_id}-discharge", "discharge_pressure", discharge, "bar"),
            ]
        )
        point_values.append(
            [
                point_id,
                "OUTPUT",
                None,
                *value_cells(f"{point_id}-power", "input_power", 2.0 + index / 10, "kW"),
            ]
        )
        point_values.append(
            [
                point_id,
                "OUTPUT",
                None,
                *value_cells(f"{point_id}-flow", "mass_flow", 160.0 + index, "kg/h"),
            ]
        )
    fixed = workbook.create_sheet("map_fixed_conditions")
    fixed.append(["map_id", *VALUE_HEADER])
    fixed.append(
        [
            "cmp-map-1",
            *value_cells(
                "map-suction-t",
                "suction_temperature",
                6.85,
                "degC",
                provided=280.0,
                provided_unit="K",
            ),
        ]
    )
    fixed.append(["cmp-map-1", *value_cells("map-frequency", "frequency", 50.0, "Hz")])

    envelopes = workbook.create_sheet("operating_envelopes")
    envelopes.append(
        [
            "product_id",
            "envelope_id",
            "refrigerant_label",
            "boundary_inclusive",
            "source_id",
            "automatic_selection",
            "restriction_reasons",
        ]
    )
    envelopes.append(
        [
            "synthetic-compressor-1",
            "envelope-1",
            "R744",
            True,
            "src-cmp-1",
            "ELIGIBLE",
            "[]",
        ]
    )
    envelope_axes = workbook.create_sheet("envelope_axes")
    envelope_axes.append(
        [
            "envelope_id",
            "axis_id",
            "order",
            "canonical_name",
            "native_name",
            "role",
            "physical_kind",
            "original_unit",
            "canonical_unit",
            "source_id",
        ]
    )
    for axis_id, order, name, native in (
        ("env-suction", 0, "suction_pressure", "Suction pressure"),
        ("env-discharge", 1, "discharge_pressure", "Discharge pressure"),
    ):
        envelope_axes.append(
            [
                "envelope-1",
                axis_id,
                order,
                name,
                native,
                "CANONICAL",
                "PRESSURE",
                "bar",
                "Pa",
                "src-cmp-1",
            ]
        )
    envelope_vertices = workbook.create_sheet("envelope_vertices")
    envelope_vertices.append(
        [
            "envelope_id",
            "vertex_order",
            "axis_id",
            "original_value",
            "original_unit",
            "provided_value_si",
            "provided_canonical_unit",
            "source_id",
        ]
    )
    for order, (suction, discharge) in enumerate(
        ((30.0, 80.0), (40.0, 80.0), (40.0, 90.0), (30.0, 90.0))
    ):
        envelope_vertices.append(
            ["envelope-1", order, "env-suction", suction, "bar", None, None, "src-cmp-1"]
        )
        envelope_vertices.append(
            [
                "envelope-1",
                order,
                "env-discharge",
                discharge,
                "bar",
                None,
                None,
                "src-cmp-1",
            ]
        )
    envelope_fixed = workbook.create_sheet("envelope_fixed_conditions")
    envelope_fixed.append(["envelope_id", *VALUE_HEADER])
    envelope_fixed.append(
        [
            "envelope-1",
            *value_cells("env-suction-t", "suction_temperature", 6.85, "degC"),
        ]
    )
    envelope_fixed.append(["envelope-1", *value_cells("env-frequency", "frequency", 50.0, "Hz")])


def write_detailed_workbook(path: Path, *, incomplete_map: bool = False) -> None:
    write_workbook(path, refrigerants='["R744"]')
    workbook = load_workbook(path)
    add_detailed_compressor_sheets(workbook, incomplete_map=incomplete_map)
    workbook.save(path)
    workbook.close()


def _column(sheet: Any, name: str) -> int:
    headers = list(next(sheet.iter_rows(values_only=True)))
    return headers.index(name) + 1


def convert_detailed_workbook_to_valve(path: Path) -> None:
    workbook = load_workbook(path)
    workbook["products"].cell(
        row=2, column=_column(workbook["products"], "component_type")
    ).value = "expansion_valve"
    workbook["rated_points"].delete_rows(2, workbook["rated_points"].max_row)
    workbook["rated_point_values"].delete_rows(2, workbook["rated_point_values"].max_row)

    maps = workbook["performance_maps"]
    maps.cell(row=2, column=_column(maps, "map_kind")).value = "EXPANSION_VALVE"
    axes = workbook["map_axes"]
    for row, axis_id, name, native in (
        (2, "axis-inlet", "inlet_pressure", "Inlet pressure"),
        (3, "axis-outlet", "outlet_pressure", "Outlet pressure"),
    ):
        axes.cell(row=row, column=_column(axes, "axis_id")).value = axis_id
        axes.cell(row=row, column=_column(axes, "canonical_name")).value = name
        axes.cell(row=row, column=_column(axes, "native_name")).value = native

    point_values = workbook["performance_point_values"]
    role_col = _column(point_values, "value_role")
    axis_col = _column(point_values, "axis_id")
    canonical_col = _column(point_values, "canonical_name")
    original_col = _column(point_values, "original_value")
    for row in range(point_values.max_row, 1, -1):
        if point_values.cell(row=row, column=canonical_col).value == "input_power":
            point_values.delete_rows(row)
            continue
        if point_values.cell(row=row, column=role_col).value != "COORDINATE":
            continue
        if point_values.cell(row=row, column=axis_col).value == "axis-suction":
            point_values.cell(row=row, column=axis_col).value = "axis-inlet"
            point_values.cell(row=row, column=canonical_col).value = "inlet_pressure"
            point_values.cell(row=row, column=original_col).value += 60.0
        else:
            point_values.cell(row=row, column=axis_col).value = "axis-outlet"
            point_values.cell(row=row, column=canonical_col).value = "outlet_pressure"
            point_values.cell(row=row, column=original_col).value -= 50.0

    fixed = workbook["map_fixed_conditions"]
    fixed.cell(row=2, column=_column(fixed, "canonical_name")).value = "inlet_temperature"
    fixed.cell(row=3, column=_column(fixed, "canonical_name")).value = "opening"
    fixed.cell(row=3, column=_column(fixed, "original_value")).value = 0.5
    fixed.cell(row=3, column=_column(fixed, "original_unit")).value = "dimensionless"

    envelope_axes = workbook["envelope_axes"]
    for row, axis_id, name, native in (
        (2, "env-inlet", "inlet_pressure", "Inlet pressure"),
        (3, "env-outlet", "outlet_pressure", "Outlet pressure"),
    ):
        envelope_axes.cell(row=row, column=_column(envelope_axes, "axis_id")).value = axis_id
        envelope_axes.cell(row=row, column=_column(envelope_axes, "canonical_name")).value = name
        envelope_axes.cell(row=row, column=_column(envelope_axes, "native_name")).value = native
    vertices = workbook["envelope_vertices"]
    vertex_axis_col = _column(vertices, "axis_id")
    vertex_value_col = _column(vertices, "original_value")
    for row in range(2, vertices.max_row + 1):
        if vertices.cell(row=row, column=vertex_axis_col).value == "env-suction":
            vertices.cell(row=row, column=vertex_axis_col).value = "env-inlet"
            vertices.cell(row=row, column=vertex_value_col).value += 60.0
        else:
            vertices.cell(row=row, column=vertex_axis_col).value = "env-outlet"
            vertices.cell(row=row, column=vertex_value_col).value -= 50.0
    envelope_fixed = workbook["envelope_fixed_conditions"]
    envelope_fixed.cell(
        row=2, column=_column(envelope_fixed, "canonical_name")
    ).value = "inlet_temperature"
    envelope_fixed.cell(row=3, column=_column(envelope_fixed, "canonical_name")).value = "opening"
    envelope_fixed.cell(row=3, column=_column(envelope_fixed, "original_value")).value = 0.5
    envelope_fixed.cell(
        row=3, column=_column(envelope_fixed, "original_unit")
    ).value = "dimensionless"
    workbook.save(path)
    workbook.close()


def add_derived_source(
    workbook: Workbook,
    *,
    source_id: str = "src-derived-1",
    input_refs: str = '["rated-power"]',
    condition_source_id: str = "src-cmp-1",
    with_condition: bool = True,
) -> None:
    workbook["data_sources"].append(
        [
            "synthetic-compressor-1",
            source_id,
            "PHYSICS_DERIVED",
            None,
            None,
            None,
            None,
            f"{source_id}-method",
            "1.0.0",
            f"{source_id}-equation",
            input_refs,
            "[]",
        ]
    )
    if not with_condition:
        return
    if "derived_source_conditions" not in workbook.sheetnames:
        conditions = workbook.create_sheet("derived_source_conditions")
        conditions.append(["derived_source_id", *VALUE_HEADER])
    conditions = workbook["derived_source_conditions"]
    conditions.append(
        [
            source_id,
            *value_cells(
                f"{source_id}-condition-pressure",
                "suction_pressure",
                30.0,
                "bar",
                source_id=condition_source_id,
            ),
        ]
    )


def write_workbook(
    path: Path,
    *,
    product_id: str = "synthetic-compressor-1",
    component_type: str = "compressor",
    refrigerants: str = '["R744", "R134a"]',
    capacity: object = 2.5,
    capacity_unit: object = "kW",
) -> None:
    """Write an explicitly synthetic workbook; it is not manufacturer evidence."""
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    products = workbook.active
    products.title = "products"
    products.append(PRODUCT_HEADER)
    products.append(
        [
            product_id,
            component_type,
            "Synthetic Manufacturer",
            "Mock Model",
            refrigerants,
            "unverified",
            True,
            capacity,
            capacity_unit,
        ]
    )
    source = workbook.create_sheet("source_metadata")
    source.append(SOURCE_HEADER)
    source.append(
        [
            product_id,
            "synthetic-test-fixture",
            "Synthetic Manufacturer",
            "Mock Model",
            "2026-09-11T00:00:00+09:00",
        ]
    )
    workbook.save(path)
    workbook.close()


def test_load_search_normalize_and_trace_exact_source(tmp_path: Path) -> None:
    path = tmp_path / "compressors" / "synthetic.xlsx"
    write_workbook(path)
    expected_hash = hashlib.sha256(path.read_bytes()).hexdigest()
    repository = ExcelComponentRepository(tmp_path)

    summary = repository.reload()
    results = repository.search(
        ComponentQuery(component_type=ComponentType.COMPRESSOR, refrigerant="r744")
    )

    assert summary.valid_files == 1
    assert summary.rejected_files == 0
    assert summary.loaded_products == 1
    assert len(results) == 1
    product = results[0]
    assert product.is_mock is True
    assert product.source.excel_file == "compressors/synthetic.xlsx"
    assert product.source.sheet == "products"
    assert product.source.row == 2
    assert product.source.file_sha256 == expected_hash
    assert product.source.document_ref == "synthetic-test-fixture"
    assert product.attributes[0].name == "rated_capacity"
    assert product.attributes[0].value.value == pytest.approx(2500.0)
    assert product.attributes[0].value.unit == "kilogram * meter ** 2 / second ** 3"


def test_discovery_order_is_relative_path_deterministic_and_ignores_temp_files(
    tmp_path: Path,
) -> None:
    write_workbook(tmp_path / "z" / "second.xlsx", product_id="z")
    write_workbook(tmp_path / "A" / "first.xlsx", product_id="a")
    write_workbook(tmp_path / "A" / "~$ignored.xlsx", product_id="ignored")
    assert [path.relative_to(tmp_path).as_posix() for path in discover_workbooks(tmp_path)] == [
        "A/first.xlsx",
        "z/second.xlsx",
    ]


@pytest.mark.parametrize(
    ("mutation", "message"),
    [
        ("missing_source_sheet", "Missing required sheet"),
        ("missing_product_column", "missing required column"),
        ("duplicate_id", "Duplicate product_id"),
        ("unpaired_unit", "require <name>_value/<name>_unit pairs"),
        ("invalid_number", "must be a finite number"),
    ],
)
def test_invalid_workbook_is_rejected(tmp_path: Path, mutation: str, message: str) -> None:
    path = tmp_path / "invalid.xlsx"
    write_workbook(path)
    workbook = load_workbook(path)
    products = workbook["products"]
    if mutation == "missing_source_sheet":
        del workbook["source_metadata"]
    elif mutation == "missing_product_column":
        products.delete_cols(1)
    elif mutation == "duplicate_id":
        products.append(list(products.iter_rows(min_row=2, max_row=2, values_only=True))[0])
    elif mutation == "unpaired_unit":
        products.cell(row=1, column=9).value = "unsupported_header"
    elif mutation == "invalid_number":
        products.cell(row=2, column=8).value = "not-a-number"
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    assert summary.valid_files == 0
    assert summary.rejected_files == 1
    assert summary.loaded_products == 0
    assert any(message in error for error in summary.errors)
    assert repository.records == ()


def test_corrupt_workbook_is_rejected(tmp_path: Path) -> None:
    path = tmp_path / "corrupt.xlsx"
    path.write_bytes(b"not an xlsx zip archive")
    summary = ExcelComponentRepository(tmp_path).reload()
    assert summary.rejected_files == 1
    assert any("Cannot open workbook" in error for error in summary.errors)


def test_duplicate_product_id_across_files_rejects_both_files(tmp_path: Path) -> None:
    write_workbook(tmp_path / "a.xlsx", product_id="duplicate")
    write_workbook(tmp_path / "b.xlsx", product_id="duplicate")
    summary = ExcelComponentRepository(tmp_path).reload()
    assert summary.valid_files == 0
    assert summary.rejected_files == 2
    assert summary.loaded_products == 0
    assert any("across files" in error for error in summary.errors)


def test_search_enforces_component_and_refrigerant_compatibility(tmp_path: Path) -> None:
    write_workbook(tmp_path / "product.xlsx")
    repository = ExcelComponentRepository(tmp_path)
    repository.reload()
    assert (
        repository.search(
            ComponentQuery(component_type=ComponentType.COMPRESSOR, refrigerant="R32")
        )
        == []
    )
    assert (
        repository.search(
            ComponentQuery(component_type=ComponentType.EVAPORATOR, refrigerant="R744")
        )
        == []
    )


def test_reload_reflects_add_modify_delete_and_is_deterministic(tmp_path: Path) -> None:
    first = tmp_path / "first.xlsx"
    second = tmp_path / "second.xlsx"
    write_workbook(first, product_id="first")
    repository = ExcelComponentRepository(tmp_path)
    initial = repository.reload()
    unchanged = repository.reload()
    assert initial.database_version == unchanged.database_version
    assert repository.records[0].product_id == "first"

    write_workbook(second, product_id="second")
    added = repository.reload()
    assert added.loaded_products == 2
    assert added.database_version != initial.database_version

    write_workbook(first, product_id="first", capacity=3.0)
    modified = repository.reload()
    assert modified.database_version != added.database_version

    first.unlink()
    deleted = repository.reload()
    assert deleted.loaded_products == 1
    assert repository.records[0].product_id == "second"


@pytest.mark.parametrize(
    ("sheet", "column", "value"),
    [
        ("products", 1, None),
        ("source_metadata", 1, None),
        ("products", 5, "R744,R134a"),
        ("products", 9, "not-a-unit"),
    ],
)
def test_invalid_file_does_not_block_valid_files(
    tmp_path: Path, sheet: str, column: int, value: object
) -> None:
    bad = tmp_path / "bad.xlsx"
    write_workbook(bad, product_id="bad")
    write_workbook(tmp_path / "good.xlsx", product_id="good")
    workbook = load_workbook(bad)
    workbook[sheet].cell(row=2, column=column).value = value
    workbook.save(bad)
    workbook.close()
    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()
    assert summary.rejected_files == 1
    assert summary.loaded_products == 1
    assert [r.product_id for r in repository.records] == ["good"]
    assert summary.errors


def test_malformed_sheet_xml_is_isolated(tmp_path: Path) -> None:
    from zipfile import ZipFile

    bad = tmp_path / "bad.xlsx"
    write_workbook(bad, product_id="bad")
    write_workbook(tmp_path / "good.xlsx", product_id="good")
    with ZipFile(bad) as archive:
        content = {name: archive.read(name) for name in archive.namelist()}
    content["xl/worksheets/sheet1.xml"] = b"<worksheet><broken"
    with ZipFile(bad, "w") as archive:
        for name, data in content.items():
            archive.writestr(name, data)
    summary = ExcelComponentRepository(tmp_path).reload()
    assert summary.rejected_files == 1
    assert summary.loaded_products == 1


def test_reload_uses_same_snapshot_for_product_and_version(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    from agent_hvac.database import index
    from agent_hvac.database.excel_loader import LOADER_VERSION, SCHEMA_VERSION

    path = tmp_path / "product.xlsx"
    write_workbook(path, product_id="original")
    original_bytes = path.read_bytes()
    original_reader = index.read_workbook_snapshot

    def read_then_remove(target: Path):
        snapshot = original_reader(target)
        target.unlink()
        return snapshot

    monkeypatch.setattr(index, "read_workbook_snapshot", read_then_remove)
    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()
    digest = hashlib.sha256(original_bytes).hexdigest()
    expected = hashlib.sha256(
        f"loader={LOADER_VERSION}\nschema={SCHEMA_VERSION}\nproduct.xlsx\0{digest}\n".encode()
    ).hexdigest()
    assert repository.records[0].product_id == "original"
    assert repository.records[0].source.file_sha256 == digest
    assert summary.database_version == f"sha256:{expected}"
    assert repository.reload().loaded_products == 0


def test_unreadable_file_is_isolated(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    from agent_hvac.database import index
    from agent_hvac.utils.exceptions import DatabaseValidationError

    write_workbook(tmp_path / "bad.xlsx", product_id="bad")
    write_workbook(tmp_path / "good.xlsx", product_id="good")
    original_reader = index.read_workbook_snapshot

    def read_with_failure(target: Path):
        if target.name == "bad.xlsx":
            raise DatabaseValidationError("Cannot read workbook: permission denied")
        return original_reader(target)

    monkeypatch.setattr(index, "read_workbook_snapshot", read_with_failure)
    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()
    assert summary.rejected_files == 1
    assert summary.loaded_products == 1
    assert summary.database_version == repository.reload().database_version


def test_uppercase_xlsx_extension_is_discovered(tmp_path: Path) -> None:
    write_workbook(tmp_path / "UPPER.XLSX")
    assert ExcelComponentRepository(tmp_path).reload().loaded_products == 1


def test_schema_020_loads_rated_map_envelope_and_normalizes_si(tmp_path: Path) -> None:
    path = tmp_path / "compressors" / "detailed.xlsx"
    write_detailed_workbook(path)
    repository = ExcelComponentRepository(tmp_path)

    summary = repository.reload()

    assert summary.valid_files == 1
    assert summary.rejected_files == 0
    assert summary.errors == ()
    product = repository.records[0]
    assert product.source.schema_version == "0.2.0"
    assert [source.source_id for source in product.data_sources] == ["src-cmp-1"]
    assert [rated.rated_point_id for rated in product.rated_points] == ["rated-1"]
    performance_map = product.performance_maps[0]
    assert performance_map.map_id == "cmp-map-1"
    assert performance_map.axes[0].original_values == (30.0, 40.0)
    assert performance_map.axes[0].values_si == (3_000_000.0, 4_000_000.0)
    assert performance_map.axes[1].values_si == (8_000_000.0, 9_000_000.0)
    assert performance_map.fixed_conditions[0].value_si.value == pytest.approx(280.0)
    assert len(performance_map.points) == 4
    envelope = product.operating_envelopes[0]
    assert envelope.envelope_id == "envelope-1"
    assert len(envelope.boundary_vertices) == 4


def test_schema_020_normalizes_basic_refrigerant_alias(tmp_path: Path) -> None:
    path = tmp_path / "alias.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    workbook["products"].cell(row=2, column=5).value = '["CO2"]'
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    assert summary.rejected_files == 0
    assert repository.records[0].supported_refrigerants == ("R744",)


def test_schema_020_rejects_unapproved_basic_refrigerant_alias(tmp_path: Path) -> None:
    path = tmp_path / "unknown-alias.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    workbook["products"].cell(row=2, column=5).value = '["R-744"]'
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    assert summary.rejected_files == 1
    assert repository.records == ()
    assert any("unapproved refrigerant alias" in error for error in summary.errors)


def test_schema_020_loads_expansion_valve_map_and_envelope(tmp_path: Path) -> None:
    path = tmp_path / "valve.xlsx"
    write_detailed_workbook(path)
    convert_detailed_workbook_to_valve(path)

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    assert summary.errors == ()
    product = repository.records[0]
    assert product.component_type == ComponentType.EXPANSION_VALVE
    assert product.rated_points == ()
    assert product.performance_maps[0].map_kind.value == "EXPANSION_VALVE"
    assert tuple(axis.canonical_name.value for axis in product.performance_maps[0].axes) == (
        "inlet_pressure",
        "outlet_pressure",
    )
    assert len(product.operating_envelopes) == 1


def test_expansion_valve_input_power_quarantines_map_only(tmp_path: Path) -> None:
    path = tmp_path / "invalid-valve.xlsx"
    write_detailed_workbook(path)
    convert_detailed_workbook_to_valve(path)
    workbook = load_workbook(path)
    point_values = workbook["performance_point_values"]
    for index in range(1, 5):
        point_values.append(
            [
                f"point-{index}",
                "OUTPUT",
                None,
                *value_cells(f"valve-power-{index}", "input_power", 1.0, "W"),
            ]
        )
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    assert repository.records[0].performance_maps == ()
    assert len(repository.records[0].operating_envelopes) == 1
    assert any("forbids input_power" in error for error in summary.errors)


def test_missing_optional_fixed_sheet_quarantines_only_dependent_map(tmp_path: Path) -> None:
    path = tmp_path / "missing-fixed.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    del workbook["map_fixed_conditions"]
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    product = repository.records[0]
    assert len(product.rated_points) == 1
    assert product.performance_maps == ()
    assert len(product.operating_envelopes) == 1
    assert any("fixed suction_temperature" in error for error in summary.errors)


def test_duplicate_value_id_quarantines_all_owning_parents(tmp_path: Path) -> None:
    path = tmp_path / "duplicate-value.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    rated_values = workbook["rated_point_values"]
    value_id_col = _column(rated_values, "value_id")
    rated_values.cell(row=2, column=value_id_col).value = "map-suction-t"
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    product = repository.records[0]
    assert product.rated_points == ()
    assert product.performance_maps == ()
    assert len(product.operating_envelopes) == 1
    assert sum("duplicate product-wide value_id" in error for error in summary.errors) == 2


def test_duplicate_source_id_across_products_invalidates_both_sources(tmp_path: Path) -> None:
    path = tmp_path / "duplicate-source.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    workbook["products"].append(
        [
            "synthetic-compressor-2",
            "compressor",
            "Synthetic Manufacturer",
            "Mock Model 2",
            '["R744"]',
            "unverified",
            True,
            2.5,
            "kW",
        ]
    )
    workbook["source_metadata"].append(
        [
            "synthetic-compressor-2",
            "synthetic-test-fixture-2",
            "Synthetic Manufacturer",
            "Mock Model 2",
            "2026-09-11T00:00:00+09:00",
        ]
    )
    workbook["data_sources"].append(
        [
            "synthetic-compressor-2",
            "src-cmp-1",
            "MANUFACTURER",
            "synthetic-detail-fixture-2",
            "Table 2",
            "Synthetic Manufacturer",
            "Mock Model 2",
            None,
            None,
            None,
            "[]",
            "[]",
        ]
    )
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    assert summary.valid_files == 1
    assert len(repository.records) == 2
    assert all(record.data_sources == () for record in repository.records)
    assert repository.records[0].rated_points == ()
    assert repository.records[0].performance_maps == ()
    assert repository.records[0].operating_envelopes == ()
    assert sum("duplicate source_id" in error for error in summary.errors) == 2


def test_derived_source_conditions_load_as_typed_si_values(tmp_path: Path) -> None:
    path = tmp_path / "derived-source.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    add_derived_source(workbook)
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    assert summary.errors == ()
    product = repository.records[0]
    assert [source.source_id for source in product.data_sources] == [
        "src-cmp-1",
        "src-derived-1",
    ]
    derived = product.data_sources[1]
    assert isinstance(derived, PhysicsDerivedDataSource)
    condition = derived.derived_method.applicable_conditions[0]
    assert condition.value_id == "src-derived-1-condition-pressure"
    assert condition.value_si.value == pytest.approx(3_000_000.0)
    assert condition.source_id == "src-cmp-1"
    assert len(product.performance_maps) == 1


def test_derived_source_without_condition_is_quarantined_only(tmp_path: Path) -> None:
    path = tmp_path / "derived-source-missing-condition.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    add_derived_source(workbook, with_condition=False)
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    product = repository.records[0]
    assert [source.source_id for source in product.data_sources] == ["src-cmp-1"]
    assert len(product.rated_points) == 1
    assert len(product.performance_maps) == 1
    assert len(product.operating_envelopes) == 1
    assert any("requires at least one applicable condition" in error for error in summary.errors)


def test_condition_row_targeting_original_source_does_not_remove_source(
    tmp_path: Path,
) -> None:
    path = tmp_path / "wrong-condition-parent.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    conditions = workbook.create_sheet("derived_source_conditions")
    conditions.append(["derived_source_id", *VALUE_HEADER])
    conditions.append(
        [
            "src-cmp-1",
            *value_cells(
                "wrong-parent-condition",
                "suction_pressure",
                30.0,
                "bar",
            ),
        ]
    )
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    product = repository.records[0]
    assert [source.source_id for source in product.data_sources] == ["src-cmp-1"]
    assert len(product.rated_points) == 1
    assert len(product.performance_maps) == 1
    assert len(product.operating_envelopes) == 1
    assert any("must reference a PHYSICS_DERIVED source" in error for error in summary.errors)


def test_unresolved_derived_condition_parent_is_orphan_only(tmp_path: Path) -> None:
    path = tmp_path / "orphan-derived-condition.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    conditions = workbook.create_sheet("derived_source_conditions")
    conditions.append(["derived_source_id", *VALUE_HEADER])
    conditions.append(
        [
            "missing-derived-source",
            *value_cells(
                "orphan-condition",
                "suction_pressure",
                30.0,
                "bar",
            ),
        ]
    )
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    product = repository.records[0]
    assert [source.source_id for source in product.data_sources] == ["src-cmp-1"]
    assert len(product.rated_points) == 1
    assert len(product.performance_maps) == 1
    assert len(product.operating_envelopes) == 1
    assert any("unresolved or duplicate derived_source_id" in error for error in summary.errors)


def test_derived_condition_may_reference_containing_source(tmp_path: Path) -> None:
    path = tmp_path / "self-sourced-derived-condition.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    add_derived_source(workbook, condition_source_id="src-derived-1")
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    assert summary.errors == ()
    derived = repository.records[0].data_sources[1]
    assert isinstance(derived, PhysicsDerivedDataSource)
    assert derived.derived_method.applicable_conditions[0].source_id == "src-derived-1"


def test_cross_product_condition_source_quarantines_derived_source_only(
    tmp_path: Path,
) -> None:
    path = tmp_path / "cross-product-derived-condition.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    workbook["products"].append(
        [
            "synthetic-compressor-2",
            "compressor",
            "Synthetic Manufacturer",
            "Mock Model 2",
            '["R744"]',
            "unverified",
            True,
            2.5,
            "kW",
        ]
    )
    workbook["source_metadata"].append(
        [
            "synthetic-compressor-2",
            "synthetic-test-fixture-2",
            "Synthetic Manufacturer",
            "Mock Model 2",
            "2026-09-11T00:00:00+09:00",
        ]
    )
    workbook["data_sources"].append(
        [
            "synthetic-compressor-2",
            "src-other-product",
            "MANUFACTURER",
            "synthetic-detail-fixture-2",
            "Table 2",
            "Synthetic Manufacturer",
            "Mock Model 2",
            None,
            None,
            None,
            "[]",
            "[]",
        ]
    )
    add_derived_source(workbook, condition_source_id="src-other-product")
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    first = next(record for record in repository.records if record.product_id.endswith("-1"))
    second = next(record for record in repository.records if record.product_id.endswith("-2"))
    assert [source.source_id for source in first.data_sources] == ["src-cmp-1"]
    assert [source.source_id for source in second.data_sources] == ["src-other-product"]
    assert len(first.performance_maps) == 1
    assert any("invalid applicable-condition source fan-out" in error for error in summary.errors)


@pytest.mark.parametrize(
    ("mutation", "message"),
    [
        ("duplicate_name", "duplicate applicable condition"),
        ("bad_unit", "incompatible with 'Pa'"),
        ("bad_audit", "provided_value_si disagrees"),
        ("nonfinite", "must be a finite number"),
    ],
)
def test_invalid_derived_condition_quarantines_only_derived_source(
    tmp_path: Path, mutation: str, message: str
) -> None:
    path = tmp_path / f"derived-condition-{mutation}.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    add_derived_source(workbook)
    conditions = workbook["derived_source_conditions"]
    if mutation == "duplicate_name":
        duplicate = list(conditions.iter_rows(min_row=2, max_row=2, values_only=True))[0]
        duplicate = list(duplicate)
        duplicate[1] = "second-condition-id"
        conditions.append(duplicate)
    elif mutation == "bad_unit":
        conditions.cell(row=2, column=_column(conditions, "original_unit")).value = "m"
    elif mutation == "bad_audit":
        conditions.cell(row=2, column=_column(conditions, "provided_value_si")).value = 1.0
        conditions.cell(row=2, column=_column(conditions, "provided_canonical_unit")).value = "Pa"
    else:
        conditions.cell(row=2, column=_column(conditions, "original_value")).value = "NaN"
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    product = repository.records[0]
    assert [source.source_id for source in product.data_sources] == ["src-cmp-1"]
    assert len(product.rated_points) == 1
    assert len(product.performance_maps) == 1
    assert len(product.operating_envelopes) == 1
    assert any(message in error for error in summary.errors)


def test_invalid_derived_source_cascades_to_downstream_source_and_parents(
    tmp_path: Path,
) -> None:
    path = tmp_path / "derived-source-cascade.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    add_derived_source(workbook, source_id="derived-upstream", with_condition=False)
    add_derived_source(
        workbook,
        source_id="derived-downstream",
        input_refs='["rated-power"]',
    )
    rated = workbook["rated_points"]
    rated.cell(row=2, column=_column(rated, "source_id")).value = "derived-upstream"
    maps = workbook["performance_maps"]
    maps.cell(row=2, column=_column(maps, "source_id")).value = "derived-downstream"
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    product = repository.records[0]
    assert [source.source_id for source in product.data_sources] == ["src-cmp-1"]
    assert product.rated_points == ()
    assert product.performance_maps == ()
    assert len(product.operating_envelopes) == 1
    assert any("requires at least one applicable condition" in error for error in summary.errors)
    assert any(
        "derived-downstream" in error and "unresolved derived input_value_refs" in error
        for error in summary.errors
    )


def test_manufacturer_source_rejects_derived_method_fields(tmp_path: Path) -> None:
    path = tmp_path / "mixed-origin-source.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    sources = workbook["data_sources"]
    sources.cell(row=2, column=_column(sources, "method_id")).value = "not-allowed"
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    product = repository.records[0]
    assert product.data_sources == ()
    assert product.rated_points == ()
    assert product.performance_maps == ()
    assert product.operating_envelopes == ()
    assert any(
        "MANUFACTURER source forbids derived method fields" in error for error in summary.errors
    )


def test_schema_010_workbook_remains_backward_compatible(tmp_path: Path) -> None:
    write_workbook(tmp_path / "legacy.xlsx")
    repository = ExcelComponentRepository(tmp_path)

    summary = repository.reload()

    assert summary.errors == ()
    product = repository.records[0]
    assert product.source.schema_version == "0.1.0"
    assert product.data_sources == ()
    assert product.rated_points == ()
    assert product.performance_maps == ()
    assert product.operating_envelopes == ()


def test_invalid_map_audit_si_quarantines_map_only(tmp_path: Path) -> None:
    path = tmp_path / "invalid-map.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    fixed = workbook["map_fixed_conditions"]
    provided_column = list(next(fixed.iter_rows(values_only=True))).index("provided_value_si") + 1
    fixed.cell(row=2, column=provided_column).value = 999.0
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    assert summary.valid_files == 1
    assert summary.rejected_files == 0
    assert summary.loaded_products == 1
    product = repository.records[0]
    assert len(product.rated_points) == 1
    assert product.performance_maps == ()
    assert len(product.operating_envelopes) == 1
    assert any(
        "map cmp-map-1" in error and "provided_value_si" in error for error in summary.errors
    )


def test_loader_parses_absolute_tolerance_with_dimension_check(tmp_path: Path) -> None:
    path = tmp_path / "tolerance.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    values = workbook["rated_point_values"]
    values.cell(row=2, column=_column(values, "tolerance_kind")).value = "ABSOLUTE"
    values.cell(row=2, column=_column(values, "tolerance_value")).value = 0.1
    values.cell(row=2, column=_column(values, "tolerance_unit")).value = "bar"
    values.cell(row=2, column=_column(values, "tolerance_reference")).value = "MEASURED_VALUE"
    values.cell(row=2, column=_column(values, "tolerance_source_id")).value = "src-cmp-1"
    values.cell(row=2, column=_column(values, "tolerance_absence_reason")).value = None
    values.cell(row=2, column=_column(values, "accuracy_validation")).value = "VERIFIABLE"
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    assert summary.errors == ()
    tolerance = repository.records[0].rated_points[0].conditions[0].tolerance.value
    assert tolerance is not None
    assert tolerance.value.value == pytest.approx(10_000.0)


def test_stray_tolerance_fields_quarantine_rated_parent_only(tmp_path: Path) -> None:
    path = tmp_path / "stray-tolerance.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    values = workbook["rated_point_values"]
    values.cell(row=2, column=_column(values, "tolerance_value")).value = 1.0
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    product = repository.records[0]
    assert product.rated_points == ()
    assert len(product.performance_maps) == 1
    assert len(product.operating_envelopes) == 1
    assert any("tolerance fields require tolerance_kind" in error for error in summary.errors)


def test_incomplete_grid_quarantines_entire_map(tmp_path: Path) -> None:
    path = tmp_path / "incomplete-map.xlsx"
    write_detailed_workbook(path, incomplete_map=True)
    repository = ExcelComponentRepository(tmp_path)

    summary = repository.reload()

    product = repository.records[0]
    assert product.performance_maps == ()
    assert len(product.rated_points) == 1
    assert len(product.operating_envelopes) == 1
    assert any("complete regular grid" in error for error in summary.errors)


def _clone_compressor_map(workbook: Workbook, suffix: str) -> str:
    map_id = f"cmp-map-{suffix}"
    maps = workbook["performance_maps"]
    parent = list(next(maps.iter_rows(min_row=2, values_only=True)))
    parent[_column(maps, "map_id") - 1] = map_id
    maps.append(parent)

    axes = workbook["map_axes"]
    for original in list(axes.iter_rows(min_row=2, values_only=True)):
        if original[_column(axes, "map_id") - 1] != "cmp-map-1":
            continue
        row = list(original)
        row[_column(axes, "map_id") - 1] = map_id
        row[_column(axes, "axis_id") - 1] += f"-{suffix}"
        axes.append(row)

    points = workbook["performance_points"]
    for original in list(points.iter_rows(min_row=2, values_only=True)):
        if original[_column(points, "map_id") - 1] != "cmp-map-1":
            continue
        row = list(original)
        row[_column(points, "map_id") - 1] = map_id
        row[_column(points, "point_id") - 1] += f"-{suffix}"
        points.append(row)

    values = workbook["performance_point_values"]
    for original in list(values.iter_rows(min_row=2, values_only=True)):
        if original[_column(values, "point_id") - 1] not in {
            "point-1",
            "point-2",
            "point-3",
            "point-4",
        }:
            continue
        row = list(original)
        row[_column(values, "point_id") - 1] += f"-{suffix}"
        row[_column(values, "value_id") - 1] += f"-{suffix}"
        if row[_column(values, "axis_id") - 1] is not None:
            row[_column(values, "axis_id") - 1] += f"-{suffix}"
        values.append(row)

    fixed = workbook["map_fixed_conditions"]
    for original in list(fixed.iter_rows(min_row=2, values_only=True)):
        if original[_column(fixed, "map_id") - 1] != "cmp-map-1":
            continue
        row = list(original)
        row[_column(fixed, "map_id") - 1] = map_id
        row[_column(fixed, "value_id") - 1] += f"-{suffix}"
        fixed.append(row)
    return map_id


@pytest.mark.parametrize("invalid_point_id", ["missing", "duplicate"])
def test_invalid_point_id_quarantines_its_entire_map(tmp_path: Path, invalid_point_id: str) -> None:
    path = tmp_path / f"{invalid_point_id}-point-id.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    points = workbook["performance_points"]
    first_invalid_row = points.max_row + 1
    if invalid_point_id == "missing":
        points.append(["cmp-map-1", None, "src-cmp-1"])
    else:
        # Both invalid rows are extra to an otherwise complete regular grid.
        points.append(["cmp-map-1", "duplicate-extra", "src-cmp-1"])
        points.append(["cmp-map-1", "duplicate-extra", "src-cmp-1"])
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    assert summary.rejected_files == 0
    assert summary.loaded_products == 1
    product = repository.records[0]
    assert product.performance_maps == ()
    assert len(product.rated_points) == 1
    assert len(product.operating_envelopes) == 1
    assert any(
        "map cmp-map-1" in error and "missing or duplicate point_id" in error
        for error in summary.errors
    )
    invalid_rows = [first_invalid_row]
    if invalid_point_id == "duplicate":
        invalid_rows.append(first_invalid_row + 1)
    for row_number in invalid_rows:
        assert any(
            f"performance_points!{row_number}:" in error and "map_id=cmp-map-1" in error
            for error in summary.errors
        )


def test_cross_map_duplicate_point_id_quarantines_both_maps_only(tmp_path: Path) -> None:
    path = tmp_path / "cross-map-point-id.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    affected_map = _clone_compressor_map(workbook, "2")
    independent_map = _clone_compressor_map(workbook, "3")
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    baseline = repository.reload()
    assert baseline.errors == ()
    assert {item.map_id for item in repository.records[0].performance_maps} == {
        "cmp-map-1",
        affected_map,
        independent_map,
    }

    workbook = load_workbook(path)
    points = workbook["performance_points"]
    first_invalid_row = points.max_row + 1
    points.append(["cmp-map-1", "cross-map-extra", "src-cmp-1"])
    points.append([affected_map, "cross-map-extra", "src-cmp-1"])
    workbook.save(path)
    workbook.close()

    summary = repository.reload()
    product = repository.records[0]
    assert summary.rejected_files == 0
    assert summary.loaded_products == 1
    assert [item.map_id for item in product.performance_maps] == [independent_map]
    assert len(product.rated_points) == 1
    assert len(product.operating_envelopes) == 1
    assert [item.source_id for item in product.data_sources] == ["src-cmp-1"]
    assert repository.search(
        ComponentQuery(component_type=ComponentType.COMPRESSOR, refrigerant="R744")
    ) == [product]
    for offset, map_id in enumerate(("cmp-map-1", affected_map)):
        assert any(
            f"performance_points!{first_invalid_row + offset}:" in error
            and f"map_id={map_id}" in error
            and "cross-map-extra" in error
            for error in summary.errors
        )
        assert any(
            f"map {map_id}" in error and "missing or duplicate point_id" in error
            for error in summary.errors
        )
    assert repository.reload().errors == summary.errors


def test_missing_point_coordinate_quarantines_map_not_workbook(tmp_path: Path) -> None:
    path = tmp_path / "missing-coordinate.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    values = workbook["performance_point_values"]
    point_col = _column(values, "point_id")
    axis_col = _column(values, "axis_id")
    for row in range(2, values.max_row + 1):
        if (
            values.cell(row=row, column=point_col).value == "point-1"
            and values.cell(row=row, column=axis_col).value == "axis-discharge"
        ):
            values.delete_rows(row)
            break
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    product = repository.records[0]
    assert summary.rejected_files == 0
    assert len(product.rated_points) == 1
    assert product.performance_maps == ()
    assert len(product.operating_envelopes) == 1
    assert any("coordinate axis IDs disagree with map" in error for error in summary.errors)


def test_missing_vertex_coordinate_quarantines_envelope_not_workbook(tmp_path: Path) -> None:
    path = tmp_path / "missing-envelope-coordinate.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    vertices = workbook["envelope_vertices"]
    order_col = _column(vertices, "vertex_order")
    axis_col = _column(vertices, "axis_id")
    for row in range(2, vertices.max_row + 1):
        if (
            vertices.cell(row=row, column=order_col).value == 0
            and vertices.cell(row=row, column=axis_col).value == "env-discharge"
        ):
            vertices.delete_rows(row)
            break
    workbook.save(path)
    workbook.close()

    repository = ExcelComponentRepository(tmp_path)
    summary = repository.reload()

    product = repository.records[0]
    assert summary.rejected_files == 0
    assert len(product.rated_points) == 1
    assert len(product.performance_maps) == 1
    assert product.operating_envelopes == ()
    assert any("coordinate axis IDs disagree with envelope" in error for error in summary.errors)


def test_missing_detail_core_sheet_keeps_static_product_with_diagnostic(tmp_path: Path) -> None:
    path = tmp_path / "missing-detail-sheet.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    del workbook["performance_point_values"]
    workbook.save(path)
    workbook.close()
    repository = ExcelComponentRepository(tmp_path)

    summary = repository.reload()

    assert summary.valid_files == 1
    assert summary.loaded_products == 1
    product = repository.records[0]
    assert product.performance_maps == ()
    assert product.rated_points == ()
    assert any("missing required detail sheet" in error for error in summary.errors)


def test_cross_product_source_quarantines_parent_not_workbook(tmp_path: Path) -> None:
    path = tmp_path / "cross-product.xlsx"
    write_detailed_workbook(path)
    workbook = load_workbook(path)
    maps = workbook["performance_maps"]
    source_column = list(next(maps.iter_rows(values_only=True))).index("source_id") + 1
    maps.cell(row=2, column=source_column).value = "src-other-product"
    workbook.save(path)
    workbook.close()
    repository = ExcelComponentRepository(tmp_path)

    summary = repository.reload()

    assert summary.valid_files == 1
    assert repository.records[0].performance_maps == ()
    assert len(repository.records[0].rated_points) == 1
    assert any("cross-product parent source_id" in error for error in summary.errors)
