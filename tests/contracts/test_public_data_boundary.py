import importlib.util
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).parents[2]
spec = importlib.util.spec_from_file_location(
    "public_data_boundary", ROOT / "scripts/check_public_data_boundary.py"
)
module = importlib.util.module_from_spec(spec)
spec.loader.exec_module(module)


def write_blank_template(root: Path) -> Path:
    path = root / module.TEMPLATE
    path.parent.mkdir(parents=True)
    workbook = Workbook()
    workbook.remove(workbook.active)
    for sheet_name in module.EXPECTED_SHEETS:
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(module.EXPECTED_HEADERS[sheet_name])
    workbook.save(path)
    workbook.close()
    return path


def test_repository_satisfies_public_data_boundary():
    assert module.check(ROOT) == []


def test_restricted_manufacturer_artifact_is_rejected(tmp_path):
    template = write_blank_template(tmp_path)
    evidence = tmp_path / "docs/product_data/manufacturer-selection.pdf"
    evidence.parent.mkdir(parents=True, exist_ok=True)
    evidence.write_bytes(b"not public")

    assert module.check(tmp_path, [template, evidence]) == [
        "Restricted public artifact: docs/product_data/manufacturer-selection.pdf"
    ]


def test_additional_workbook_is_rejected(tmp_path):
    template = write_blank_template(tmp_path)
    extra = tmp_path / "data/components/compressors/product.xlsx"
    extra.parent.mkdir(parents=True)
    extra.write_bytes(b"not public")

    assert module.check(tmp_path, [template, extra]) == [
        "Restricted public artifact: data/components/compressors/product.xlsx",
        "Public component data must remain empty: data/components/compressors/product.xlsx",
    ]


def test_populated_public_template_is_rejected(tmp_path):
    path = write_blank_template(tmp_path)
    workbook = module.load_workbook(path)
    workbook["products"]["A2"] = "synthetic-or-real-data"
    workbook.save(path)
    workbook.close()

    assert module.check(tmp_path, [path]) == ["Public template contains data: products!2"]


def test_component_data_directory_accepts_only_gitkeep(tmp_path):
    template = write_blank_template(tmp_path)
    placeholder = tmp_path / "data/components/compressors/.gitkeep"
    placeholder.parent.mkdir(parents=True)
    placeholder.touch()
    assert module.check(tmp_path, [template, placeholder]) == []

    (placeholder.parent / "record.json").write_text("{}", encoding="utf-8")
    record = placeholder.parent / "record.json"
    assert module.check(tmp_path, [template, placeholder, record]) == [
        "Public component data must remain empty: data/components/compressors/record.json"
    ]


def test_template_sheet_contract_is_enforced(tmp_path):
    path = write_blank_template(tmp_path)
    workbook = module.load_workbook(path)
    workbook.remove(workbook["envelope_fixed_conditions"])
    workbook.save(path)
    workbook.close()

    errors = module.check(tmp_path, [path])
    assert len(errors) == 1
    assert errors[0].startswith("Public template sheet contract differs:")


def test_tracked_build_artifact_is_rejected(tmp_path):
    template = write_blank_template(tmp_path)
    evidence = tmp_path / "docs/build/manufacturer-report.pdf"
    evidence.parent.mkdir(parents=True)
    evidence.write_bytes(b"not public")

    assert module.check(tmp_path, [template, evidence]) == [
        "Restricted public artifact: docs/build/manufacturer-report.pdf"
    ]


def test_modified_template_header_is_rejected(tmp_path):
    path = write_blank_template(tmp_path)
    workbook = module.load_workbook(path)
    workbook["products"]["A1"] = "arbitrary-data"
    workbook.save(path)
    workbook.close()

    errors = module.check(tmp_path, [path])
    assert len(errors) == 1
    assert errors[0].startswith("Public template header differs: products!1;")


def test_approved_gui_evidence_is_allowed(tmp_path):
    template = write_blank_template(tmp_path)
    screenshots = []
    for relative_path in sorted(module.PUBLIC_BINARY_ALLOWLIST - {module.TEMPLATE}):
        screenshot = tmp_path / relative_path
        screenshot.parent.mkdir(parents=True, exist_ok=True)
        screenshot.write_bytes(b"synthetic screenshot")
        screenshots.append(screenshot)

    assert module.check(tmp_path, [template, *screenshots]) == []
