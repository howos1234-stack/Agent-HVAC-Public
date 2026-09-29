"""Strict, deterministic ingestion for the P07 minimum Excel product schema."""

from __future__ import annotations

import hashlib
import json
import math
from dataclasses import dataclass
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Final
from xml.etree.ElementTree import ParseError
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from agent_hvac.schemas.components import ComponentType, ProductRecord, ProductSource
from agent_hvac.schemas.provenance import Parameter, Provenance, SourceType
from agent_hvac.utils.exceptions import DatabaseValidationError
from agent_hvac.utils.units import Quantity

LOADER_VERSION: Final = "0.2.0"
SCHEMA_VERSION: Final = "0.2.0"
LEGACY_SCHEMA_VERSION: Final = "0.1.0"
PRODUCTS_SHEET: Final = "products"
SOURCE_SHEET: Final = "source_metadata"

PRODUCT_COLUMNS: Final = (
    "product_id",
    "component_type",
    "manufacturer",
    "model",
    "supported_refrigerants",
    "status",
    "is_mock",
)
SOURCE_COLUMNS: Final = (
    "product_id",
    "document_ref",
    "original_manufacturer",
    "original_model",
    "retrieved_at",
)


def file_sha256(path: Path) -> str:
    """Return the digest of the exact workbook bytes used for ingestion."""
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


@dataclass(frozen=True)
class WorkbookSnapshot:
    data: bytes
    modified_at: datetime

    @property
    def sha256(self) -> str:
        return hashlib.sha256(self.data).hexdigest()


@dataclass(frozen=True)
class WorkbookLoadResult:
    records: tuple[ProductRecord, ...]
    diagnostics: tuple[str, ...]


def read_workbook_snapshot(path: Path) -> WorkbookSnapshot:
    """Use one open file for payload and metadata; parsing never reopens the path."""
    import os

    try:
        with path.open("rb") as stream:
            data = stream.read()
            modified_at = datetime.fromtimestamp(os.fstat(stream.fileno()).st_mtime, tz=UTC)
        return WorkbookSnapshot(data, modified_at)
    except OSError as exc:
        raise DatabaseValidationError(f"Cannot read workbook: {exc}") from exc


def discover_workbooks(root: Path) -> tuple[Path, ...]:
    """Discover non-temporary xlsx files in a stable relative-path order."""
    if not root.exists():
        return ()
    if not root.is_dir():
        raise DatabaseValidationError(f"Component database root is not a directory: {root}")
    return tuple(
        sorted(
            (
                path
                for path in root.rglob("*")
                if path.is_file()
                and path.suffix.casefold() == ".xlsx"
                and not path.name.startswith("~$")
            ),
            key=lambda path: (
                path.relative_to(root).as_posix().casefold(),
                path.relative_to(root).as_posix(),
            ),
        )
    )


def load_product_workbook(
    path: Path, *, database_root: Path, snapshot: WorkbookSnapshot | None = None
) -> tuple[ProductRecord, ...]:
    """Load one workbook or raise one explicit validation error for the whole file."""
    return load_product_workbook_with_diagnostics(
        path, database_root=database_root, snapshot=snapshot
    ).records


def load_product_workbook_with_diagnostics(
    path: Path, *, database_root: Path, snapshot: WorkbookSnapshot | None = None
) -> WorkbookLoadResult:
    """Load one workbook and return deterministic detailed-record diagnostics."""
    from agent_hvac.database.detailed_loader import (
        has_detail_sheets,
        load_details,
        normalize_supported_refrigerants,
    )

    snapshot = snapshot if snapshot is not None else read_workbook_snapshot(path)
    try:
        workbook = load_workbook(BytesIO(snapshot.data), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError, ValueError, KeyError, ParseError) as exc:
        raise DatabaseValidationError(f"Cannot open workbook: {exc}") from exc

    try:
        missing_sheets = {PRODUCTS_SHEET, SOURCE_SHEET}.difference(workbook.sheetnames)
        if missing_sheets:
            raise DatabaseValidationError(
                f"Missing required sheet(s): {', '.join(sorted(missing_sheets))}"
            )

        product_rows = _rows_by_header(workbook[PRODUCTS_SHEET], PRODUCT_COLUMNS)
        source_rows = _rows_by_header(workbook[SOURCE_SHEET], SOURCE_COLUMNS)
        if not product_rows:
            raise DatabaseValidationError("products sheet contains no product rows")

        sources = _unique_rows(source_rows, sheet=SOURCE_SHEET)
        products = _unique_rows(product_rows, sheet=PRODUCTS_SHEET)
        missing_sources = sorted(set(products).difference(sources))
        orphan_sources = sorted(set(sources).difference(products))
        if missing_sources:
            raise DatabaseValidationError(
                f"Missing source_metadata for product_id(s): {', '.join(missing_sources)}"
            )
        if orphan_sources:
            raise DatabaseValidationError(
                f"Orphan source_metadata product_id(s): {', '.join(orphan_sources)}"
            )

        digest = snapshot.sha256
        modified_at = snapshot.modified_at
        relative_path = path.resolve().relative_to(database_root.resolve()).as_posix()
        detailed_schema = has_detail_sheets(workbook.sheetnames)
        records = tuple(
            _build_record(
                product_row=products[product_id],
                source_row=sources[product_id],
                relative_path=relative_path,
                digest=digest,
                modified_at=modified_at,
                schema_version=SCHEMA_VERSION if detailed_schema else LEGACY_SCHEMA_VERSION,
            )
            for product_id in sorted(products)
        )
        if not detailed_schema:
            return WorkbookLoadResult(records=records, diagnostics=())
        records = normalize_supported_refrigerants(records)
        details = load_details(
            workbook,
            records,
            relative_path=relative_path,
            digest=digest,
        )
        return WorkbookLoadResult(records=details.records, diagnostics=details.diagnostics)
    except (ValueError, TypeError, KeyError, OSError, BadZipFile, ParseError) as exc:
        raise DatabaseValidationError(f"Invalid workbook: {exc}") from exc
    finally:
        workbook.close()


def _rows_by_header(sheet: Any, required: tuple[str, ...]) -> list[tuple[int, dict[str, object]]]:
    rows = sheet.iter_rows(values_only=True)
    try:
        raw_header = next(rows)
    except StopIteration as exc:
        raise DatabaseValidationError(f"{sheet.title} sheet is empty") from exc
    header = [_clean_header(value) for value in raw_header]
    if len(header) != len(set(header)):
        raise DatabaseValidationError(f"{sheet.title} contains duplicate column names")
    missing = sorted(set(required).difference(header))
    if missing:
        raise DatabaseValidationError(
            f"{sheet.title} missing required column(s): {', '.join(missing)}"
        )
    result: list[tuple[int, dict[str, object]]] = []
    for row_number, values in enumerate(rows, start=2):
        if all(value is None or str(value).strip() == "" for value in values):
            continue
        result.append((row_number, dict(zip(header, values, strict=False))))
    return result


def _clean_header(value: object) -> str:
    if not isinstance(value, str) or not value.strip():
        raise DatabaseValidationError("Every populated column must have a non-empty text header")
    return value.strip()


def _unique_rows(
    rows: list[tuple[int, dict[str, object]]], *, sheet: str
) -> dict[str, tuple[int, dict[str, object]]]:
    result: dict[str, tuple[int, dict[str, object]]] = {}
    for row_number, row in rows:
        product_id = _required_text(row.get("product_id"), f"{sheet}!A{row_number} product_id")
        if product_id in result:
            first_row = result[product_id][0]
            raise DatabaseValidationError(
                f"Duplicate product_id {product_id!r} in {sheet} rows {first_row} and {row_number}"
            )
        result[product_id] = (row_number, row)
    return result


def _build_record(
    *,
    product_row: tuple[int, dict[str, object]],
    source_row: tuple[int, dict[str, object]],
    relative_path: str,
    digest: str,
    modified_at: datetime,
    schema_version: str,
) -> ProductRecord:
    row_number, product = product_row
    _, source = source_row
    product_id = _required_text(product["product_id"], "product_id")
    source_ref = f"{relative_path}#{PRODUCTS_SHEET}!{row_number}"
    try:
        return ProductRecord(
            product_id=product_id,
            component_type=ComponentType(
                _required_text(product["component_type"], "component_type")
            ),
            manufacturer=_required_text(product["manufacturer"], "manufacturer"),
            model=_required_text(product["model"], "model"),
            supported_refrigerants=_refrigerants(product["supported_refrigerants"]),
            source=ProductSource(
                excel_file=relative_path,
                sheet=PRODUCTS_SHEET,
                row=row_number,
                document_ref=_required_text(source["document_ref"], "document_ref"),
                original_manufacturer=_required_text(
                    source["original_manufacturer"], "original_manufacturer"
                ),
                original_model=_required_text(source["original_model"], "original_model"),
                workbook_modified_at=modified_at,
                retrieved_at=_aware_datetime(source["retrieved_at"]),
                file_sha256=digest,
                loader_version=LOADER_VERSION,
                schema_version=schema_version,
            ),
            status=_required_text(product["status"], "status"),
            attributes=_attribute_parameters(product, source_ref=source_ref),
            is_mock=_boolean(product["is_mock"], "is_mock"),
        )
    except (ValidationError, ValueError, TypeError) as exc:
        raise DatabaseValidationError(f"Invalid products row {row_number}: {exc}") from exc


def _attribute_parameters(row: dict[str, object], *, source_ref: str) -> tuple[Parameter, ...]:
    reserved = set(PRODUCT_COLUMNS)
    value_columns = {name.removesuffix("_value") for name in row if name.endswith("_value")}
    unit_columns = {name.removesuffix("_unit") for name in row if name.endswith("_unit")}
    unpaired = sorted(value_columns.symmetric_difference(unit_columns))
    if unpaired:
        raise DatabaseValidationError(
            "Physical attribute columns require <name>_value/<name>_unit pairs: "
            f"{', '.join(unpaired)}"
        )
    unknown = sorted(
        name
        for name, value in row.items()
        if name not in reserved
        and not name.endswith(("_value", "_unit"))
        and value is not None
        and str(value).strip()
    )
    if unknown:
        raise DatabaseValidationError(
            f"Unsupported populated product column(s): {', '.join(unknown)}"
        )

    parameters: list[Parameter] = []
    for name in sorted(value_columns):
        value = row.get(f"{name}_value")
        unit = row.get(f"{name}_unit")
        if value is None and unit is None:
            continue
        if isinstance(value, bool) or not isinstance(value, (int, float)):
            raise DatabaseValidationError(f"{name}_value must be a finite number")
        if not math.isfinite(float(value)):
            raise DatabaseValidationError(f"{name}_value must be a finite number")
        parameters.append(
            Parameter(
                name=name,
                value=Quantity(value=float(value), unit=_required_text(unit, f"{name}_unit")),
                provenance=Provenance(source_type=SourceType.PRODUCT_DB, source_ref=source_ref),
            )
        )
    return tuple(parameters)


def _refrigerants(value: object) -> tuple[str, ...]:
    if not isinstance(value, str):
        raise ValueError("supported_refrigerants must be a JSON array string")
    try:
        decoded = json.loads(value)
    except json.JSONDecodeError as exc:
        raise ValueError("supported_refrigerants must be a JSON array string") from exc
    if not isinstance(decoded, list) or not decoded:
        raise ValueError("supported_refrigerants must be a non-empty JSON array")
    values = tuple(_required_text(item, "supported_refrigerants item") for item in decoded)
    if len({item.casefold() for item in values}) != len(values):
        raise ValueError("supported_refrigerants contains duplicates")
    return values


def _required_text(value: object, field: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{field} must be non-empty text")
    return value.strip()


def _boolean(value: object, field: str) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, str) and value.strip().casefold() in {"true", "false"}:
        return value.strip().casefold() == "true"
    raise ValueError(f"{field} must be TRUE or FALSE")


def _aware_datetime(value: object) -> datetime:
    if isinstance(value, datetime):
        parsed = value
    elif isinstance(value, str):
        try:
            parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        except ValueError as exc:
            raise ValueError("retrieved_at must be an ISO-8601 datetime") from exc
    else:
        raise ValueError("retrieved_at must be an ISO-8601 datetime")
    if parsed.tzinfo is None or parsed.utcoffset() is None:
        raise ValueError("retrieved_at must include a timezone")
    return parsed
